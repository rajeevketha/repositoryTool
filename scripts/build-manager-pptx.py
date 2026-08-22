#!/usr/bin/env python3
"""Build the manager briefing PowerPoint (16:9) for Google Slides import."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

BG = RGBColor(0x12, 0x15, 0x1C)
CARD = RGBColor(0x1C, 0x21, 0x2C)
TEXT = RGBColor(0xE6, 0xE4, 0xDE)
MUTED = RGBColor(0x8F, 0x96, 0xA3)
COPPER = RGBColor(0xC5, 0xBB, 0xA8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0x3A, 0x3F, 0x4A)

W = Inches(13.333)
H = Inches(7.5)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_bg(slide):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    set_fill(box, BG)
    sp = box._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)


def add_text(slide, left, top, width, height, text, size=18, color=TEXT, bold=False, name=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = name or "Calibri"
    return tf


def kicker(slide, text):
    add_text(slide, Inches(0.7), Inches(0.38), Inches(12), Inches(0.35), text.upper(), 12, COPPER, True)


def title(slide, text, size=32):
    add_text(slide, Inches(0.7), Inches(0.72), Inches(12), Inches(1.15), text, size, WHITE, True)


def body(slide, text, top=2.0, size=20):
    add_text(slide, Inches(0.7), Inches(top), Inches(11.9), Inches(4.6), text, size, TEXT)


def card(slide, left, top, width, height, heading, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    set_fill(shape, CARD)
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    p.text = heading
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Calibri"
    q = tf.add_paragraph()
    q.text = text
    q.font.size = Pt(13)
    q.font.color.rgb = MUTED
    q.font.name = "Calibri"
    q.space_before = Pt(8)


def table_slide(slide, headers, rows):
    rows_n = 1 + len(rows)
    cols = len(headers)
    table = slide.shapes.add_table(rows_n, cols, Inches(0.7), Inches(2.05), Inches(11.9), Inches(4.6)).table
    for i in range(cols):
        table.columns[i].width = Inches(2.2 if i == 0 else 4.85)
    def paint(cell, text, header=False):
        cell.text = text
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(12 if not header else 11)
            p.font.bold = header or cell == table.cell(0, 0)
            p.font.color.rgb = COPPER if header else TEXT
            p.font.name = "Calibri"
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD
    for c, h in enumerate(headers):
        paint(table.cell(0, c), h, True)
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            paint(table.cell(r, c), val)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Manager briefing")
    title(s, "Promote Salesforce configuration\nwithout a developer toolchain", 36)
    body(
        s,
        "OrgFlow is a Chrome extension for admins and configurators. They pick the change they just built in a sandbox and send that same package to QA, then production.\n\n"
        "No new IDE. No mandatory Git repo. No Copado-class license. A guided From → To hop in the browser they already use.",
        3.1,
        20,
    )

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Why this is on a manager’s desk")
    title(s, "Small config changes are still expensive")
    card(s, Inches(0.7), Inches(2.1), Inches(3.8), Inches(4.3), "Handoffs", "A field or layout often waits on someone who knows VS Code, a change set, or a release tool. The work took twenty minutes. The promotion takes a day.")
    card(s, Inches(4.7), Inches(2.1), Inches(3.8), Inches(4.3), "Tool sprawl", "Workbench for zips. VS Code for source. A DevOps platform for releases. None match “send this named change to the next org.”")
    card(s, Inches(8.7), Inches(2.1), Inches(3.8), Inches(4.3), "Wrong package risk", "People rebuild the list in each org. OrgFlow keeps one snapshot and walks it DEC → QA → prod so both receive the same files.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Fit")
    title(s, "Who should use it — and who should not")
    card(s, Inches(0.7), Inches(2.1), Inches(5.8), Inches(4.4), "Good fit", "Admins and configurators who live in Setup.\nTeams that promote a handful of named items at a time.\nOrgs that want a shared history later, not on day one.\nManagers who want a pilot without a platform RFP.")
    card(s, Inches(6.8), Inches(2.1), Inches(5.8), Inches(4.4), "Not a replacement for", "Developer source work (Apex, LWC, CI) in VS Code.\nEnterprise release trains (Copado, Gearset, Flosum).\nData / record migration.\nAutomatic dependency analysis of a whole org.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "The work")
    title(s, "One path. Four steps. They cannot skip.")
    card(s, Inches(0.55), Inches(2.1), Inches(2.95), Inches(4.4), "1  Start", "Log into both Salesforce orgs in Chrome. Detect them. Set From and To. They must be different.")
    card(s, Inches(3.65), Inches(2.1), Inches(2.95), Inches(4.4), "2  Package", "OrgFlow shows what just changed in From. They tick the items to move — not a blank XML file.")
    card(s, Inches(6.75), Inches(2.1), Inches(2.95), Inches(4.4), "3  Retrieve", "Pull those items from From. Optional Jira. Same files can travel to the next org later.")
    card(s, Inches(9.85), Inches(2.1), Inches(2.95), Inches(4.4), "4  Confirm", "See the path and what will land in To. Validate first if you want. Then Deploy.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Adoption")
    title(s, "Training is measured in minutes, not weeks")
    card(s, Inches(0.7), Inches(2.1), Inches(5.8), Inches(4.4), "What people already have", "Chrome. Two Salesforce logins. That is the setup. OrgFlow uses the existing browser session. No Connected App, no CLI install, no project folder.")
    card(s, Inches(6.8), Inches(2.1), Inches(5.8), Inches(4.4), "What the product refuses to do", "Next stays off until From and To are different orgs. Deploy stays off until retrieve succeeded. Changing From clears the old selection so yesterday’s package does not silently go to the wrong place.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Versus Workbench")
    title(s, "Workbench is a console. OrgFlow is a promotion.")
    table_slide(
        s,
        ["", "Salesforce Workbench", "OrgFlow"],
        [
            ["Role", "Expert API tool. Retrieve or deploy a zip someone already prepared.", "Daily path for “I built this in sandbox — send it to QA.”"],
            ["Skill", "You must know package.xml and Metadata API layout.", "Tick names. The package is built for them."],
            ["Repeatability", "Each org is a new session. Easy to send a different list next time.", "One snapshot. Same files can hop to the next environment."],
            ["Oversight", "No From → To bar. Easy to deploy back into the same org.", "From and To must differ. Confirm shows the path before send."],
        ],
    )

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Versus VS Code")
    title(s, "Do not send configurators into an IDE to move a layout")
    table_slide(
        s,
        ["", "VS Code + Salesforce CLI", "OrgFlow"],
        [
            ["Staffing", "Developers and technical admins. Scratch orgs, source, CI.", "Configurators who work in Setup all day."],
            ["Onboarding", "Authorize an org, create a DX project, learn retrieve/deploy.", "Install a Chrome extension. Detect orgs. Follow Next."],
            ["System of record", "The Git repo is the product. Correct for engineering.", "Git is optional. A shared repo is a later decision, not a gate."],
            ["Small change", "Worth it for Apex and LWC. Overhead for one field and a layout.", "Built for that small change."],
        ],
    )

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Versus tools you may already pay for")
    title(s, "Complement, do not rip and replace")
    card(s, Inches(0.7), Inches(2.05), Inches(5.8), Inches(2.1), "Change Sets", "Slow, easy to miss a dependency, awkward to reuse. OrgFlow is a named list, one hop, result on the same screen.")
    card(s, Inches(6.8), Inches(2.05), Inches(5.8), Inches(2.1), "Copado / Gearset / Flosum", "Right for release trains and compliance programs. Wrong as the only way to send two fields to QA. OrgFlow sits beside those platforms.")
    card(s, Inches(0.7), Inches(4.35), Inches(5.8), Inches(2.1), "Git-only process", "A repo without a simple picker still requires a developer. OrgFlow can write versions into GitHub, GitLab, or Azure after members are chosen.")
    card(s, Inches(6.8), Inches(4.35), Inches(5.8), Inches(2.1), "Explicit limits", "No data (records). No full-org compare. No automatic dependency graph. People promote what they selected, not what a scanner guessed.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Controls")
    title(s, "What you can tell security and release owners")
    card(s, Inches(0.7), Inches(2.1), Inches(3.8), Inches(4.3), "Access", "Uses the Salesforce session already in Chrome. No Connected App to approve. Git tokens, if used, stay in that browser.")
    card(s, Inches(4.7), Inches(2.1), Inches(3.8), Inches(4.3), "Blast radius", "Named members only. From and To cannot be the same org. Validate is a dry run. Production can still run Apex tests.")
    card(s, Inches(8.7), Inches(2.1), Inches(3.8), Inches(4.3), "Audit", "On this device: local snapshots, optional Jira. With a release repo: Jira key and comment are required; versions live in the team warehouse.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "Suggested next step")
    title(s, "A two-week pilot, not a program")
    card(s, Inches(0.7), Inches(2.1), Inches(5.8), Inches(4.4), "Pilot", "Two people, two sandbox or disposable orgs.\nOne real config change (field + layout or a flow).\nPromote sandbox → QA with storage on this device.\nOptional: turn on a test Git repo in week two.")
    card(s, Inches(6.8), Inches(2.1), Inches(5.8), Inches(4.4), "Success looks like", "They complete a hop without a developer.\nQA receives the same members they picked.\nNo new platform license to evaluate first.\nYou decide later if the team needs a shared repo.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    kicker(s, "How to get it")
    title(s, "Install, detect, send", 36)
    body(
        s,
        "Chrome Web Store listing, or load the unpacked extension from the team zip. Two logged-in Salesforce orgs. From and To must be different.\n\n"
        "https://chromewebstore.google.com/detail/orgflow/cldfmnjjonlakccebbfnaaihneonhflc\n\n"
        "This file opens in Google Slides: Drive → New → File upload → Open with Google Slides.",
        3.15,
        20,
    )

    out = "/workspace/docs/orgflow-manager-briefing.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
